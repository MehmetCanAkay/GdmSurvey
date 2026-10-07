"""
Excel soru havuzunu questions.json dosyasına dönüştürür.

Soru metni, tip, eksen ve alt-tema her çalıştırmada Excel'den yenilenir.
Soru başına puanlanacak CAS maddeleri CAS madde haritasından okunur.
Aynı soru kimliği için daha önce elle girilmiş checklist maddeleri korunur.
"""

import json
from pathlib import Path

from openpyxl import load_workbook

from src.domain import (
    CAS_ALWAYS_SCORED,
    CAS_ITEMS,
    CAS_MAP_CODES,
    EXPECTED_GUIDELINE_QUESTION_COUNT,
    EXPECTED_PATIENT_QUESTION_COUNT,
    EXPECTED_QUESTION_COUNT,
    QuestionType,
    axis_from_excel,
    question_type_from_excel,
)

EXCEL_FILENAME = "GDM_Final_30_Soru_Birlesik.xlsx"
CAS_MAP_FILENAME = "GDM_CAS_Madde_Haritasi.xlsx"
QUESTION_SHEET_MARK = "SORU NO"
SUMMARY_AXIS_HEADER = "EKSEN"
CAS_MAP_SHEET = "CAS Haritası"
CAS_SCORED_MARK = "✓"
CAS_SKIPPED_MARKS = {"–", "-", ""}


class QuestionImportError(ValueError):
    """Excel içeriği çalışma kurallarına uymadığında fırlatılır."""


def import_questions(excel_path: Path, json_path: Path, cas_map_path: Path | None = None) -> dict:
    """
    Soru Excel'ini ve CAS madde haritasını okur, doğrular ve questions.json yazar.

    CAS haritası verilmezse soru Excel'iyle aynı klasördeki dosya kullanılır.
    Dönüş değeri yazılan JSON gövdesidir.
    """
    excel_path = Path(excel_path)
    json_path = Path(json_path)
    cas_map_path = Path(cas_map_path) if cas_map_path else excel_path.parent / CAS_MAP_FILENAME
    if not excel_path.exists():
        raise QuestionImportError(f"Excel dosyası bulunamadı: {excel_path}")
    if not cas_map_path.exists():
        raise QuestionImportError(f"CAS madde haritası bulunamadı: {cas_map_path}")

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        questions = _read_questions(workbook)
        summary = _read_summary(workbook)
    finally:
        workbook.close()
    cas_map = read_cas_map(cas_map_path)

    _validate(questions, summary)
    _validate_cas_map(questions, cas_map)
    preserved = _existing_checklists(json_path)
    for question in questions:
        question["cas_items"] = cas_map[question["id"]]
        question["checklist"] = preserved.get(question["id"], [])

    payload = {"source": excel_path.name, "cas_source": cas_map_path.name, "questions": questions}
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return payload


def _read_questions(workbook) -> list[dict]:
    """Soru havuzu sayfasını başlık adına göre okur."""
    sheet = _find_question_sheet(workbook)
    rows = list(sheet.iter_rows(values_only=True))
    header_index = _header_index(rows, QUESTION_SHEET_MARK)
    header = [_cell_text(value) for value in rows[header_index]]
    columns = _column_map(header)

    questions = []
    seen: set[str] = set()
    problems: list[str] = []
    for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if _row_empty(row):
            continue
        question_id = _cell_text(row[columns["id"]])
        if not question_id or question_id.casefold() == "toplam":
            continue
        if question_id in seen:
            problems.append(f"Satır {offset}: tekrarlanan soru kimliği {question_id}")
            continue
        seen.add(question_id)
        text = _cell_text(row[columns["text"]])
        raw_type = _cell_text(row[columns["type"]])
        raw_axis = _cell_text(row[columns["axis"]])
        sub_theme = _cell_text(row[columns["sub_theme"]])
        if not text:
            problems.append(f"{question_id}: soru metni boş")
            continue
        try:
            question_type = question_type_from_excel(raw_type)
            axis = axis_from_excel(raw_axis)
        except ValueError as exc:
            problems.append(f"{question_id}: {exc}")
            continue
        questions.append(
            {
                "id": question_id,
                "type": question_type.value,
                "axis": axis.value,
                "sub_theme": sub_theme,
                "text": text,
            }
        )

    if problems:
        raise QuestionImportError("\n".join(problems))
    return questions


def _read_summary(workbook) -> dict[str, dict]:
    """Özet sayfasındaki eksen toplamlarını ve soru listelerini okur."""
    sheet = _find_summary_sheet(workbook)
    rows = list(sheet.iter_rows(values_only=True))
    header_index = None
    for index, row in enumerate(rows):
        if any(_cell_text(cell).casefold() == SUMMARY_AXIS_HEADER.casefold() for cell in row):
            header_index = index
            break
    if header_index is None:
        raise QuestionImportError("Özet sayfasında eksen başlığı bulunamadı.")

    header = [_cell_text(value) for value in rows[header_index]]
    axis_col = _find_column(header, "EKSEN")
    total_col = _find_column(header, "TOPLAM")
    ids_col = _find_column(header, "SORU NUMARALARI")

    summary: dict[str, dict] = {}
    for row in rows[header_index + 1 :]:
        label = _cell_text(row[axis_col]) if axis_col < len(row) else ""
        if not label or label.casefold() == "toplam":
            continue
        try:
            axis = axis_from_excel(label)
        except ValueError as exc:
            raise QuestionImportError(str(exc)) from exc
        total = row[total_col] if total_col < len(row) else None
        id_text = _cell_text(row[ids_col]) if ids_col < len(row) else ""
        question_ids = [part.strip() for part in id_text.split(",") if part.strip()]
        summary[axis.value] = {"total": int(total), "ids": question_ids}
    return summary


def _validate(questions: list[dict], summary: dict[str, dict]) -> None:
    """Sayı, tip dağılımı ve özet sayfası ile havuzu karşılaştırır."""
    problems: list[str] = []
    if len(questions) != EXPECTED_QUESTION_COUNT:
        problems.append(f"Soru sayısı {len(questions)}; beklenen {EXPECTED_QUESTION_COUNT}.")

    patient = sum(1 for q in questions if q["type"] == QuestionType.HASTA_TUREVLI.value)
    guideline = sum(1 for q in questions if q["type"] == QuestionType.KILAVUZ_VAKA.value)
    if patient != EXPECTED_PATIENT_QUESTION_COUNT:
        problems.append(
            f"Hasta-türevli soru sayısı {patient}; beklenen {EXPECTED_PATIENT_QUESTION_COUNT}."
        )
    if guideline != EXPECTED_GUIDELINE_QUESTION_COUNT:
        problems.append(
            f"Kılavuz vakası sayısı {guideline}; beklenen {EXPECTED_GUIDELINE_QUESTION_COUNT}."
        )

    actual_ids: dict[str, list[str]] = {}
    for question in questions:
        actual_ids.setdefault(question["axis"], []).append(question["id"])

    for axis, expected in summary.items():
        found = actual_ids.get(axis, [])
        if len(found) != expected["total"]:
            problems.append(
                f"{axis} ekseninde {len(found)} soru var; özette {expected['total']}."
            )
        if expected["ids"] and found != expected["ids"]:
            problems.append(
                f"{axis} soru sırası özetteki listeyle uyuşmuyor: {found} / {expected['ids']}"
            )

    missing_axes = sorted(set(actual_ids) - set(summary))
    if missing_axes:
        problems.append(f"Özet sayfasında olmayan eksenler: {', '.join(missing_axes)}")

    if problems:
        raise QuestionImportError("\n".join(problems))


def read_cas_map(cas_map_path: Path) -> dict[str, list[str]]:
    """
    CAS madde haritasını soru kimliği -> puanlanan CAS anahtarları sözlüğüne çevirir.

    Hücrede ✓ puanlanır, – puanlanmaz demektir. Başka değer hata sayılır.
    Anahtarlar CAS_ITEMS sırasıyla döner.
    """
    workbook = load_workbook(cas_map_path, read_only=True, data_only=True)
    try:
        if CAS_MAP_SHEET not in workbook.sheetnames:
            raise QuestionImportError(f"CAS haritasında '{CAS_MAP_SHEET}' sayfası yok.")
        rows = list(workbook[CAS_MAP_SHEET].iter_rows(values_only=True))
    finally:
        workbook.close()

    header_index = None
    for index, row in enumerate(rows):
        if row and _cell_text(row[0]) == "Soru":
            header_index = index
            break
    if header_index is None:
        raise QuestionImportError("CAS haritasında 'Soru' başlığı bulunamadı.")
    header = [_cell_text(value) for value in rows[header_index]]
    code_columns = {}
    for code in CAS_MAP_CODES:
        if code not in header:
            raise QuestionImportError(f"CAS haritasında {code} sütunu yok.")
        code_columns[code] = header.index(code)

    order = [item.key for item in CAS_ITEMS]
    mapping: dict[str, list[str]] = {}
    problems: list[str] = []
    for row in rows[header_index + 1 :]:
        question_id = _cell_text(row[0]) if row else ""
        if not question_id or question_id.casefold() == "toplam":
            continue
        if question_id in mapping:
            problems.append(f"CAS haritasında tekrarlanan soru: {question_id}")
            continue
        keys = []
        for code, column in code_columns.items():
            mark = _cell_text(row[column]) if column < len(row) else ""
            if mark == CAS_SCORED_MARK:
                keys.append(CAS_MAP_CODES[code])
            elif mark not in CAS_SKIPPED_MARKS:
                problems.append(f"CAS haritası {question_id}/{code}: tanınmayan değer '{mark}'")
        mapping[question_id] = sorted(keys, key=order.index)
    if problems:
        raise QuestionImportError("\n".join(problems))
    return mapping


def _validate_cas_map(questions: list[dict], cas_map: dict[str, list[str]]) -> None:
    """CAS haritasının soru havuzuyla aynı kimlikleri taşıdığını ve kültürel varsayım maddesini her soruda içerdiğini denetler."""
    problems: list[str] = []
    question_ids = [question["id"] for question in questions]
    missing = [question_id for question_id in question_ids if question_id not in cas_map]
    extra = sorted(set(cas_map) - set(question_ids))
    if missing:
        problems.append(f"CAS haritasında olmayan sorular: {', '.join(missing)}")
    if extra:
        problems.append(f"Soru havuzunda olmayan CAS satırları: {', '.join(extra)}")
    for question_id in question_ids:
        if question_id in cas_map and CAS_ALWAYS_SCORED not in cas_map[question_id]:
            problems.append(f"{question_id}: kültürel varsayım maddesi (M5) her soruda puanlanmalı.")
    if problems:
        raise QuestionImportError("\n".join(problems))


def _existing_checklists(json_path: Path) -> dict[str, list]:
    """Var olan JSON dosyasındaki geçerli checklist listelerini okur."""
    if not json_path.exists():
        return {}
    raw = json.loads(json_path.read_text(encoding="utf-8"))
    if isinstance(raw, list):
        # Eski düz liste biçiminde checklist yoktur.
        return {}
    if not isinstance(raw, dict):
        return {}

    preserved: dict[str, list] = {}
    for question in raw.get("questions", []):
        checklist = question.get("checklist") or []
        if not checklist:
            continue
        _validate_checklist(str(question.get("id")), checklist)
        preserved[str(question["id"])] = checklist
    return preserved


def _validate_checklist(question_id: str, checklist: list) -> None:
    """Checklist maddelerinin id, etiket ve kategori taşıdığını denetler."""
    seen: set[str] = set()
    for index, item in enumerate(checklist, start=1):
        if not isinstance(item, dict):
            raise QuestionImportError(f"{question_id} checklist maddesi {index} nesne değil.")
        for field in ("id", "label", "category"):
            value = item.get(field)
            if not isinstance(value, str) or not value.strip():
                raise QuestionImportError(
                    f"{question_id} checklist maddesi {index} için '{field}' boş."
                )
        item_id = item["id"].strip()
        if item_id in seen:
            raise QuestionImportError(f"{question_id} checklist içinde tekrarlanan id: {item_id}")
        seen.add(item_id)


def _find_question_sheet(workbook):
    """İçinde SORU NO başlığı olan sayfayı bulur."""
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            if any(_cell_text(cell) == QUESTION_SHEET_MARK for cell in row):
                return sheet
    raise QuestionImportError("Soru sayfasında 'SORU NO' başlığı bulunamadı.")


def _find_summary_sheet(workbook):
    """Eksen dağılımını içeren özet sayfasını bulur."""
    for sheet in workbook.worksheets:
        title = sheet.title.casefold()
        if "özet" in title or "ozet" in title:
            return sheet
    raise QuestionImportError("Özet istatistik sayfası bulunamadı.")


def _header_index(rows: list, marker: str) -> int:
    """Başlık satırının indeksini döndürür."""
    for index, row in enumerate(rows):
        if any(_cell_text(cell) == marker for cell in row):
            return index
    raise QuestionImportError(f"Başlık bulunamadı: {marker}")


def _column_map(header: list[str]) -> dict[str, int]:
    """Soru sayfası sütunlarını başlık metnine göre eşler."""
    return {
        "id": _find_column(header, "SORU NO"),
        "type": _find_column(header, "TİP"),
        "axis": _find_column(header, "EKSEN"),
        "sub_theme": _find_column(header, "ALT-TEMA"),
        "text": _find_column(header, "SORU METNİ"),
    }


def _find_column(header: list[str], label: str) -> int:
    """Başlık listesinde etiketi içeren ilk sütunun indeksini döndürür."""
    needle = label.casefold()
    for index, name in enumerate(header):
        if needle in name.casefold():
            return index
    raise QuestionImportError(f"Sütun bulunamadı: {label}")


def _cell_text(value) -> str:
    """Hücre değerini kırpılmış metne çevirir."""
    if value is None:
        return ""
    return str(value).strip()


def _row_empty(row) -> bool:
    """Satırdaki tüm hücreler boşsa True döner."""
    return all(not _cell_text(cell) for cell in row)


def main() -> None:
    """Komut satırından soru Excel'ini ve CAS haritasını questions.json dosyasına dönüştürür."""
    root = Path(__file__).resolve().parent.parent
    payload = import_questions(
        root / "data" / EXCEL_FILENAME,
        root / "data" / "questions.json",
    )
    print(f"{len(payload['questions'])} soru yazıldı: data/questions.json")


if __name__ == "__main__":
    main()
